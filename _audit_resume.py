import docx
p = r"C:/Users/yangweihua/Desktop/简历/output/resume-ai-20260829/stage3/杨伟华-AI应用开发工程师-简历.docx"
doc = docx.Document(p)
kw = ["AI", "智能", "Agent", "ReAct", "工具", "RAG", "临床", "知识库", "Text2SQL", "SQL",
      "轮转", "排班", "考勤", "通知", "RBAC", "租户", "MCP", "评测", "流式", "SSE",
      "FastAPI", "微服务", "Python", "Laravel", "接口", "向量", "嵌入", "检索"]
print("==== 简历含 AI/技术关键词的段落 ====")
for i, t in enumerate(doc.paragraphs):
    s = t.text.strip()
    if s and any(k in s for k in kw):
        print(f"[{i}] {s}")
print("\n==== 表格单元格 ====")
for ti, tbl in enumerate(doc.tables):
    for ri, row in enumerate(tbl.rows):
        for c in row.cells:
            s = c.text.strip()
            if s and any(k in s for k in kw):
                print(f"  T{ti}R{ri}: {s[:140]}")
